"""Genera docs/Documento_Diseno_Ajustado.docx a partir de los resultados reales (logs/ y docs/*.json)."""
import csv, inspect, json, sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from creditscore.acciones import Actuadores
from creditscore.agente import AgenteCreditScore
from creditscore.decision import ModuloDecision
from creditscore.percepcion import ModuloPercepcion

R = Path(__file__).resolve().parents[1]
URL_REPO = "https://github.com/stevencontre/creditscore-agent"
URL_APP = "https://creditscore-agent-7pyg7dzx7yiykdggfwoinw.streamlit.app/"   # pegar aqui la URL publica https://<nombre>.streamlit.app tras desplegar
D, L = R / "docs", R / "logs"
MZ = json.loads((D / "resultados_matriz.json").read_text(encoding="utf-8"))
MX = json.loads((D / "matriz_decision.json").read_text(encoding="utf-8"))
PR = json.loads((L / "resultados_pruebas.json").read_text(encoding="utf-8"))
ES = json.loads((L / "escenarios.json").read_text(encoding="utf-8"))
with open(L / "resultados_pruebas.csv", encoding="utf-8-sig", newline="") as f:
    TABLA = list(csv.reader(f, delimiter=";"))
f2 = lambda x, n=2: f"{x:.{n}f}".replace(".", ",")
pct = lambda x, n=1: f"{x*100:.{n}f} %".replace(".", ",")

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for a in ("left_margin", "right_margin"): setattr(sec, a, Cm(2.2))
sec.top_margin = sec.bottom_margin = Cm(2.2)
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
st.paragraph_format.space_after = Pt(5); st.paragraph_format.line_spacing = 1.12
for n, sz in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 10.5)):
    s = doc.styles[n]; s.font.name = "Arial"; s.font.size = Pt(sz); s.font.bold = True; s.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
    s.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    s.paragraph_format.space_before = Pt(12 if n == "Heading 1" else 8); s.paragraph_format.keep_with_next = True


def H(t, n=1): return doc.add_heading(t, n)


def P(t="", bold=False, italic=False, align=None, size=None):
    p = doc.add_paragraph(); r = p.add_run(t); r.bold, r.italic = bold, italic
    if size: r.font.size = Pt(size)
    if align == "c": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def PB(partes):
    """Parrafo con negritas: lista de (texto, negrita)."""
    p = doc.add_paragraph()
    for t, b in partes:
        p.add_run(t).bold = b
    return p


def B(t, bold_pref=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_pref:
        p.add_run(bold_pref).bold = True
    p.add_run(t); p.paragraph_format.space_after = Pt(2); return p


def sombrear(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), color); tcPr.append(sh)


def T(cab, filas, anchos, fs=8, cap=None):
    if cap:
        c = P(cap, bold=True, size=8.5); c.paragraph_format.keep_with_next = True; c.paragraph_format.space_after = Pt(2)
    t = doc.add_table(rows=1, cols=len(cab)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(cab):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(fs)
        sombrear(c, "D9E2F3")
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for fila in filas:
        cells = t.add_row().cells
        for i, v in enumerate(fila):
            cells[i].text = ""; par = cells[i].paragraphs[0]; par.paragraph_format.space_after = Pt(0)
            txt = str(v); neg = txt.startswith("**")
            r = par.add_run(txt.strip("*")); r.font.size = Pt(fs); r.bold = neg
    for i, w in enumerate(anchos):
        t.columns[i].width = Cm(w)
    for row in t.rows:
        for i, w in enumerate(anchos):
            row.cells[i].width = Cm(w)
        trPr = row._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def FIG(path, ancho, pie, salto=False):
    doc.add_picture(str(path), width=Cm(ancho)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.page_break_before = salto
    p = P(pie, italic=True, align="c", size=8.5); p.paragraph_format.space_after = Pt(8)


def FIGFIT(path, max_w, max_h, pie, salto=False):
    from PIL import Image
    w, h = Image.open(path).size
    ancho = min(max_w, max_h * w / h)
    FIG(path, ancho, pie, salto)


def CODE(txt, fs=7):
    for linea in txt.rstrip().splitlines():
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.left_indent = Cm(0.4)
        r = p.add_run(linea if linea else " "); r.font.name = "Consolas"; r.font.size = Pt(fs)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        pPr = p._p.get_or_add_pPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "F2F2F2"); pPr.append(sh)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def src(obj, desde=None, hasta=None, max_l=40):
    lineas = inspect.getsource(obj).splitlines()
    ini = next((i for i, l in enumerate(lineas) if desde and desde in l), 0)
    fin = next((i for i in range(ini + 1, len(lineas)) if hasta and hasta in lineas[i]), len(lineas)) if hasta else len(lineas)
    corte = lineas[ini:fin][:max_l]
    sangria = min(len(l) - len(l.lstrip()) for l in corte if l.strip())
    return "\n".join(l[sangria:] for l in corte)


def pagina_landscape(on):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    if on:
        s.orientation = WD_ORIENT.LANDSCAPE; s.page_width, s.page_height = Cm(29.7), Cm(21)
    else:
        s.orientation = WD_ORIENT.PORTRAIT; s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(1.8 if on else 2.2); s.top_margin = s.bottom_margin = Cm(2.0)


# pie de pagina con numero
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.add_run("Página ").font.size = Pt(8)
for tipo in ("begin", "instr", "end"):
    r = fp.add_run(); r.font.size = Pt(8)
    if tipo == "instr":
        e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = " PAGE "
    else:
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tipo)
    r._r.append(e)

# ============================ PORTADA ============================
for t, sz, b in (("UNIVERSIDAD TECNOLÓGICA DE BOLÍVAR", 12, True), ("INGENIERÍA DE SISTEMAS Y COMPUTACIÓN", 11, True)):
    P(t, bold=b, align="c", size=sz)
P("")
P("Diseño de un Agente Inteligente — Documento ajustado", bold=True, align="c", size=16)
P("CREDITSCORE-AGENT — Agente inteligente de evaluación de riesgo crediticio", bold=True, align="c", size=13)
P("Versión 2: incluye selección justificada, desarrollo del agente y pruebas de desempeño", italic=True, align="c")
P("")
P("Presentado por:", bold=True, align="c")
P("Steven David Contreras — T00077010", align="c")
P("Andres Camilo Castaño — T00077278", align="c")
P("Curso: Inteligencia Artificial", align="c")
P("Docente: Edwin Alexander Puertas", align="c")
P("Cartagena, 17/10/2026", align="c")

# ============================ TABLA DE CAMBIOS ============================
H("Tabla de cambios respecto de la entrega anterior (Diseño inicial, 27/08/2026)", 1)
P("Cada fila indica la sección afectada, lo que decía el documento anterior, lo que dice ahora y el motivo. Los cambios responden a la retroalimentación del docente (texto con marcas de herramientas de citación, encabezados que no correspondían al contenido y ausencia de verificación) y a los nuevos requisitos de la actividad.")
T(["Sección", "Contenido anterior", "Contenido nuevo", "Justificación"], [
    ["Portada", "Un solo integrante (Steven David Contreras).", "Dos integrantes de la pareja.", "La entrega es por parejas."],
    ["Todo el documento", "Texto con marcas literales de herramientas de búsqueda y citación (por ejemplo «cite,turn0search4», «filecite turn0file1 L22-L24») en al menos cinco secciones.", "Marcas eliminadas; el texto se redactó sin ellas y las fuentes se citan por autor y año en Referencias.", "Retroalimentación del docente: eran artefactos no revisados."],
    ["Secciones 4 a 7", "Encabezados desordenados: las entradas aparecían bajo «Descripción del entorno», las salidas bajo «Entradas y sensores», el texto SMART bajo «Salidas y actuadores» y «Medidas de desempeño» estaba vacío.", "Cada encabezado contiene su contenido: 4 entorno; 5 entradas y sensores (Tabla 1); 6 salidas y actuadores (Tabla 2); 7 medidas de desempeño.", "Retroalimentación del docente: contenido pegado sin reordenar."],
    ["7. Medidas de desempeño", "Dos metas SMART (AUC-ROC ≥ 0,75 y reducción del 20 % del tiempo) y métricas complementarias sin umbral.", "Se conservan las metas y se agregan fórmulas, umbrales fijados antes de ejecutar y la reformulación de la medida de tiempo (sección 9.7).", "La reducción del 20 % no puede probarse sin una línea base manual; se reformuló y se documentó el cambio."],
    ["Sección A (nueva)", "No existía.", "Selección justificada del agente: matriz ponderada, cálculo, sensibilidad y justificación.", "Requisito de la Fase 1 de la actividad."],
    ["9. Desarrollo del agente (nueva)", "No existía; las conclusiones eran la sección 9.", "Nueve subsecciones: arquitectura, tecnologías, entradas, salidas, UML, diferencias diseño–implementación, pruebas, limitaciones y enlaces.", "Requisito de la Fase 2. Las conclusiones pasan a ser la sección 10."],
    ["10. Conclusiones", "Concluía viabilidad técnica sin evidencia de ejecución.", "Conclusiones actualizadas con resultados reales y sus límites (datos sintéticos).", "Coherencia con las pruebas ejecutadas."],
    ["Referencias", "Russell y Norvig (2021) aparecía dos veces; sin verificación de enlaces.", "Referencia duplicada fusionada; se agregan las fuentes técnicas usadas. Enlaces marcados para verificación.", "Calidad de las referencias."],
], [2.6, 5.0, 5.0, 3.8], fs=7.5)

# ============================ SECCION A: MATRIZ ============================
H("A. Selección justificada del agente (sección nueva)", 1)
P("La actividad pide implementar uno de los dos agentes diseñados en la entrega anterior: CREDITSCORE-AGENT (riesgo crediticio, diseño de Steven David Contreras) y Cartrescue-agent (recuperación de carritos abandonados, diseño del compañero de pareja). Los criterios y pesos se definieron antes de puntuar, con los valores sugeridos por la actividad.")
PB([("Advertencia de método. ", True), ("El documento de diseño de Cartrescue-agent no estaba disponible al elaborar esta versión. Por eso todas las puntuaciones del agente B son estimaciones basadas solo en su nombre y propósito general, y están marcadas como ESTIMACIÓN. Las del agente A se apoyan en su diseño y en el prototipo construido, lo que introduce un sesgo de información a favor de A. Los integrantes deben revisar las filas de B contra su diseño real y volver a ejecutar «python scripts/matriz_decision.py» después de editar docs/matriz_decision.json.", False)])
A1 = []
for c in MX["criterios"]:
    A1.append([c["nombre"], f"{c['peso']*100:.0f} %", str(c["A"]["p"]), c["A"]["evidencia"], str(c["B"]["p"]), c["B"]["evidencia"]])
A1.append(["**Puntaje ponderado S**", "**100 %**", f"**{f2(MZ['S_A'])}**", "", f"**{f2(MZ['S_B'])}**", ""])
T(["Criterio", "Peso (w)", "Agente A: puntaje (1–5)", "Agente A: evidencia", "Agente B: puntaje (1–5)", "Agente B: evidencia"], A1, [2.3, 1.1, 1.3, 5.2, 1.3, 5.2], fs=7, cap="Tabla A1. Matriz de decisión ponderada (A = CREDITSCORE-AGENT; B = Cartrescue-agent, puntajes estimados)")
H("A.1 Cálculo del puntaje ponderado", 2)
P("Se aplica S(agente) = Σ (wᵢ × pᵢ), con Σ wᵢ = 1 y puntaje máximo 5.")
w = [c["peso"] for c in MX["criterios"]]
pa = [c["A"]["p"] for c in MX["criterios"]]; pb = [c["B"]["p"] for c in MX["criterios"]]
CODE("S(A) = " + " + ".join(f"{f2(x,2)}×{p}" for x, p in zip(w, pa)) + f" = {f2(MZ['S_A'])}\nS(B) = " + " + ".join(f"{f2(x,2)}×{p}" for x, p in zip(w, pb)) + f" = {f2(MZ['S_B'])}", fs=8)
H("A.2 Análisis de sensibilidad", 2)
P(f"El criterio más influyente se definió como el que más aporta a la diferencia entre agentes, es decir, el de mayor wᵢ·|pA − pB|. Resultó ser «{MZ['criterio_mas_influyente']}» (aporte de {f2(max(MZ['contribucion_a_la_diferencia'].values()))} puntos). Su peso se modificó en +10 y −10 puntos porcentuales y los demás pesos se reescalaron proporcionalmente para sumar 100 %.")
nom = [c["nombre"].split(" ")[0] + " " + (c["nombre"].split(" ")[1] if len(c["nombre"].split(" ")) > 1 else "") for c in MX["criterios"]]
fil = []
for k, e in MZ["escenarios"].items():
    fil.append([k.capitalize(), " / ".join(f"{x*100:.1f}".replace(".", ",") for x in e["pesos"]), f2(e["S_A"], 3), f2(e["S_B"], 3), "A" if e["elegido"] == "A" else e["elegido"]])
rb = MZ["robustez_B_mas_1_en_estimaciones"]
fil.append(["Robustez: puntajes estimados de B + 1 (tope 5), pesos originales", "25 / 20 / 20 / 15 / 10 / 10", f2(rb["S_A"], 3), f2(rb["S_B"], 3), rb["elegido"]])
T(["Escenario", "Pesos en % (viabilidad / datos / medibilidad / complejidad / costo / valor)", "S(A)", "S(B)", "Elegido"], fil, [4.4, 6.4, 1.5, 1.5, 1.6], fs=8, cap="Tabla A2. Sensibilidad de la decisión")
P(f"Resultado: con los puntajes actuales la decisión se mantiene en A en los cuatro escenarios de pesos (original, +10 pp, −10 pp y pesos iguales). La ventaja mínima de A es de {f2(MZ['escenarios']['-10 pp']['S_A']-MZ['escenarios']['-10 pp']['S_B'],2)} puntos (escenario −10 pp). En cambio, la decisión no es robusta a los puntajes de B: si el diseño real de Cartrescue-agent justificara un punto más en cada criterio estimado, B obtendría {f2(rb['S_B'])} frente a {f2(rb['S_A'])} y pasaría a ser la mejor opción. La ventaja de A tiene que confirmarse con el diseño real de B.")
H("A.3 Justificación de la selección (una página)", 2)
P(f"Se selecciona CREDITSCORE-AGENT con un puntaje ponderado de {f2(MZ['S_A'])} sobre 5, frente a {f2(MZ['S_B'])} estimado para Cartrescue-agent. La diferencia proviene sobre todo de dos criterios: complejidad de IA (A aporta un agente de aprendizaje con decisión por utilidad y abstención por incertidumbre, mientras que B se estimó como agente de reglas o utilidad) y costo y recursos (A funciona con Python y numpy sin APIs de pago, mientras que el envío real de recordatorios de B suele requerir servicios con cuota).")
P("Resultado de la sensibilidad: al subir o bajar 10 puntos el peso de la complejidad de IA, A conserva el primer lugar (4,16 frente a 3,49 y 3,94 frente a 3,61). Con pesos iguales A también gana (4,17 frente a 3,50). La decisión depende de los puntajes estimados de B, como se indicó en A.2.")
P("Ventajas de la elección: medibilidad directa mediante AUC-ROC, calibración, trazabilidad y latencia; una arquitectura que cubre los siete componentes IBM; y ejecución completa en un portátil, sin costo. Limitaciones: no hay datos reales de solicitudes, por lo que se usa un simulador sintético y los resultados demuestran el funcionamiento de la tubería, no el desempeño en una entidad financiera; la medida de reducción del 20 % del tiempo exige una línea base manual que aún no existe; y la evaluación de equidad entre subgrupos no es posible con estos datos.")
P("Qué se pierde al descartar a Cartrescue-agent: un problema con retroalimentación inmediata (el cliente compra o no tras el recordatorio), que permitiría medir el aprendizaje en ciclos cortos, y un dominio con menor riesgo ético que las decisiones crediticias. Si el diseño real de Cartrescue-agent incluye aprendizaje y datos disponibles, esta decisión debe revisarse.")

# ============================ SECCIONES 1-8 (DISEÑO ORIGINAL CORREGIDO) ============================
H("1. Introducción y planteamiento del problema", 1)
P("El proceso de evaluación crediticia es una actividad crítica para las entidades financieras, porque requiere estimar la probabilidad de incumplimiento de las obligaciones adquiridas por un solicitante. En un procedimiento predominantemente manual, el análisis de información financiera, historial de pagos, nivel de endeudamiento y líneas de crédito abiertas puede demandar tiempo considerable y presentar variaciones derivadas de la carga operativa, los criterios aplicados y la disponibilidad de información. El diseño de un agente inteligente permite automatizar parte del proceso de percepción, análisis y recomendación, manteniendo supervisión humana para los casos con incertidumbre o características atípicas.")
P("El CREDITSCORE-AGENT evalúa el riesgo de incumplimiento de nuevas solicitudes de crédito. Según el caso de uso del catálogo del curso, el objetivo es alcanzar un AUC-ROC mínimo de 0,75 en la predicción de incumplimiento y reducir en 20 % el tiempo de respuesta frente al proceso manual, usando una muestra de 5.000 solicitudes históricas. El agente recibe variables financieras del solicitante, estima una probabilidad de incumplimiento y una calificación de riesgo, contrasta el resultado con decisiones históricas sobre solicitudes similares y registra la evaluación en el sistema de originación. El resultado final de cada solicitud se conserva para recalibrar periódicamente el modelo (Curso de Inteligencia Artificial, 2026).")
H("2. Tipo de agente inteligente", 1)
P("Según la clasificación de Russell y Norvig (2021), el CREDITSCORE-AGENT es principalmente un agente de aprendizaje, con características complementarias de un agente basado en utilidad. Es de aprendizaje porque usa resultados históricos para construir y actualizar el modelo de predicción: cada solicitud tiene variables de entrada y un resultado observado posteriormente, lo que permite evaluar las predicciones y recalibrar. Es también basado en utilidad porque la decisión no es una condición binaria: pondera el riesgo estimado, la incertidumbre, los criterios de la política crediticia y la conveniencia de remitir casos a revisión humana.")
H("3. Tipos de ambiente", 1)
T(["Propiedad", "Clasificación", "Justificación"], [
    ["Observabilidad", "Parcialmente observable", "Solo se dispone de las variables suministradas y del historial; no se observan los factores futuros que afectan la capacidad de pago."],
    ["Determinismo", "Estocástico", "La relación entre las características y el incumplimiento es probabilística."],
    ["Episodicidad", "Secuencial", "Cada solicitud es una unidad de evaluación, pero las decisiones previas y sus resultados influyen en la recalibración y en evaluaciones futuras."],
    ["Dinamismo", "Dinámico", "Cambian las condiciones económicas, las políticas y la composición de los datos; el agente debe recalibrarse."],
    ["Discretización", "Principalmente continuo", "Ingresos, relación deuda-ingreso y probabilidades son continuos; las acciones (aprobar, rechazar, remitir) son discretas."],
    ["Número de agentes", "Agente único con humanos y sistemas externos", "El núcleo es un agente que interactúa con analistas y con el sistema de originación."],
], [3.2, 4.2, 9.0], fs=8.5)
H("4. Descripción del entorno", 1)
P("El entorno está compuesto por la entidad financiera, los solicitantes, los analistas de revisión, el sistema de originación, las fuentes de información financiera y la infraestructura de datos para entrenamiento, validación y seguimiento del modelo. El agente se ubica entre la recepción de una solicitud y la generación de una recomendación de riesgo.")
P("En condiciones normales, el sistema recibe una solicitud con información suficiente, valida la integridad de los datos, ejecuta el modelo, genera una probabilidad de incumplimiento y asigna una categoría de riesgo. Después registra el resultado y puede recomendar aprobación, rechazo o revisión humana según los umbrales y políticas de la organización. En condiciones excepcionales (datos incompletos, valores inconsistentes, resultados fuera de la distribución histórica o incertidumbre elevada) el agente debe abstenerse de producir una decisión automática y remitir el caso a revisión especializada.")
H("5. Entradas y sensores", 1)
P("Los sensores del agente son componentes de software que capturan, validan y transforman las señales del entorno: API, conectores con bases de datos, formularios digitales, módulos de validación y procesos de preprocesamiento. Su función es convertir las percepciones en una representación estructurada apta para el modelo predictivo. El caso de uso contempla ingresos, historial de pagos, relación deuda-ingreso, líneas de crédito abiertas y el historial de decisiones sobre solicitudes similares.")
T(["Entrada", "Tipo de dato", "Captura"], [
    ["Ingresos", "Numérico continuo", "Formulario, sistema financiero o API de verificación"],
    ["Historial de pagos", "Datos históricos tabulares; variables categóricas y temporales", "Base de datos o API de información crediticia"],
    ["Relación deuda-ingreso", "Numérico continuo", "Cálculo automático a partir de ingresos y obligaciones registradas"],
    ["Líneas de crédito abiertas", "Numérico discreto y atributos tabulares", "Base de datos o API de información crediticia"],
    ["Historial de decisiones", "Tabular y categórico", "Base de datos del sistema de originación"],
    ["Resultado observado del crédito", "Categórico binario y temporal", "Sistema de seguimiento de cartera"],
], [4.2, 6.0, 6.2], fs=8.5, cap="Tabla 1. Entradas, tipos de datos y mecanismos de captura")
H("6. Salidas y actuadores", 1)
P("Las salidas son las acciones o recomendaciones del agente. Los actuadores son los mecanismos mediante los cuales modifican el estado del sistema de originación o activan la intervención humana.")
T(["Salida", "Actuador", "Efecto sobre el entorno"], [
    ["Aprobación recomendada", "API del sistema de originación", "Registra la recomendación y permite continuar el flujo de aprobación según la política institucional."],
    ["Rechazo recomendado", "API del sistema de originación", "Registra el resultado y detiene o deriva el flujo de formalización del crédito según las reglas vigentes."],
    ["Remisión a revisión humana", "Módulo de notificaciones y bandeja de analistas", "Asigna el caso a un responsable humano para una evaluación adicional."],
    ["Calificación de riesgo", "Módulo de scoring y registro", "Incorpora la puntuación al expediente crediticio y permite aplicar los umbrales definidos."],
    ["Registro histórico", "Base de datos de auditoría y aprendizaje", "Conserva entradas, salida, versión del modelo y resultado posterior para trazabilidad y recalibración."],
], [4.0, 5.0, 7.4], fs=8.5, cap="Tabla 2. Salidas, actuadores y efecto sobre el entorno")
H("7. Medidas de desempeño (SMART)", 1)
B(" evaluar automáticamente el riesgo de incumplimiento de solicitudes de crédito a partir de variables financieras e históricas y generar una recomendación de riesgo.", "Específica:")
B(" alcanzar un AUC-ROC mínimo de 0,75 en la predicción de incumplimiento y reducir en 20 % el tiempo de respuesta respecto del proceso manual.", "Medible:")
B(" validar el desempeño con una muestra de 5.000 solicitudes históricas, como establece el caso de uso.", "Alcanzable:")
B(" mejorar la rapidez y consistencia de la evaluación, dando apoyo analítico a la decisión crediticia.", "Relevante:")
B(" comprobar los indicadores durante una fase piloto definida por la organización y antes de una eventual puesta en producción.", "Temporal:")
P("Como métricas complementarias se recomienda monitorizar precisión, recall de la clase de incumplimiento, calibración de probabilidades, tasa de falsos positivos, estabilidad temporal y diferencias de desempeño entre subgrupos. En la sección 9.7 cada medida se convierte en una prueba con fórmula, umbral previo y resultado numérico; la reducción del 20 % se reformuló por falta de línea base.")
H("8. Componentes de la arquitectura de IA agéntica (marco IBM)", 1)
P("IBM (2025) describe los agentes de IA como sistemas con componentes interconectados que permiten percibir, procesar, razonar, actuar, comunicarse y aprender. La plantilla organiza el diseño en siete componentes: percepción y manejo de entradas; planificación y descomposición de tareas; memoria; razonamiento y toma de decisiones; llamada de herramientas y acción; comunicación; y aprendizaje y adaptación.")
for tit, txt in [
    ("Percepción y manejo de entradas", "Recibe la solicitud y transforma los datos en una representación estructurada: validación de campos, detección de valores ausentes o inconsistentes, normalización y transformaciones del modelo, además de verificar que la solicitud esté dentro del dominio esperado. Una interpretación incorrecta de los datos afecta las decisiones posteriores."),
    ("Planificación y descomposición de tareas", "Descompone la evaluación en: recibir la solicitud, validar la información, consultar los datos, ejecutar la predicción, contrastar con casos similares, verificar reglas, determinar la necesidad de revisión humana y registrar la salida."),
    ("Memoria", "La memoria de corto plazo conserva el contexto de la solicitud en evaluación (variables, transformaciones, resultado del modelo, reglas activadas y acciones). La de largo plazo almacena historial de solicitudes, resultados observados, decisiones, versiones del modelo y métricas, con controles de acceso."),
    ("Razonamiento y toma de decisiones", "Combina el resultado probabilístico con las políticas y umbrales de la entidad, la calidad de los datos, la similitud con casos históricos y la incertidumbre, para decidir si la solicitud sigue un flujo automatizado o se remite a un analista. No sustituye la gobernanza institucional."),
    ("Llamada de herramientas y acción", "Usa herramientas de software (por API o funciones estructuradas) para consultar fuentes de datos, ejecutar el modelo, obtener información del sistema de originación y registrar el resultado."),
    ("Comunicación", "Intercambia información con sistemas institucionales y analistas. Para el analista produce una salida estructurada con calificación, probabilidad, factores relevantes y razón de la remisión; para los sistemas, interfaces estructuradas."),
    ("Aprendizaje y adaptación", "Usa los resultados observados como retroalimentación: compara probabilidades estimadas con resultados reales, detecta cambios en la distribución de los datos y recalibra el modelo cuando se cumplen los criterios establecidos.")]:
    PB([(tit + ". ", True), (txt, False)])

# ============================ SECCION 9: DESARROLLO ============================
H("9. Desarrollo del agente", 1)
P("Esta sección describe el agente implementado en el repositorio (carpeta creditscore_agent). Todas las cifras proceden de ejecuciones reales registradas en la carpeta logs/. Los datos de solicitudes son sintéticos, generados por un simulador propio.")
H("9.1 Arquitectura", 2)
FIGFIT(D / "diagramas" / "arquitectura.png", 16.4, 9.0, "Figura 1. Diagrama de componentes del CREDITSCORE-AGENT (fuente: docs/diagramas/arquitectura.dot)")
P("El ciclo Percepción → Decisión → Acción está en AgenteCreditScore.evaluar(). Los sensores (formulario y buró) entregan datos al módulo de percepción, que valida y construye el vector de características. El módulo de decisión consulta el modelo predictivo y la memoria de largo plazo (casos similares), aplica las reglas y la función de utilidad, y entrega una Decision. Los actuadores ejecutan esa decisión sobre el entorno (originación, bandeja de analistas, expediente y auditoría) y el registro de eventos deja constancia con marca de tiempo. El resultado observado del crédito llega después desde la cartera y alimenta la revisión de aprendizaje.")
T(["Componente de la arquitectura", "Archivo / clase", "Qué hace"], [
    ["Sensores y percepción", "percepcion.py · ModuloPercepcion", "Lee formulario y buró (con reintentos), valida tipos y rangos, calcula el DTI y construye 5 características."],
    ["Modelo predictivo", "modelo.py · ModeloRiesgo", "Regresión logística con regularización L2 (Newton-Raphson ponderado); contribuciones por variable para explicar."],
    ["Estado interno (memoria)", "memoria.py · MemoriaCortoPlazo, MemoriaLargoPlazo", "Contexto de la solicitud; historial con resultados y búsqueda de los 25 casos más similares."],
    ["Módulo de decisión", "decision.py · ModuloDecision", "Reglas R1 a R4 y minimización del costo esperado entre aprobar, rechazar y remitir."],
    ["Actuadores y acción", "acciones.py · Actuadores", "Originación con reintentos y outbox, bandeja de analistas, expediente de calificación y auditoría."],
    ["Planificación y ciclo", "agente.py · AgenteCreditScore", "Plan de ocho tareas y orquestación; manejo de errores con decisión segura (revisión humana)."],
    ["Aprendizaje y adaptación", "agente.py · revisar_aprendizaje()", "Detecta deriva (PSI y brecha de calibración) y recalibra el modelo con ponderación de casos recientes."],
    ["Entorno simulado", "entorno.py, simulador.py, fabrica.py", "Genera solicitudes sintéticas, simula buró y originación con fallos inyectables y entrega el resultado observado."],
    ["Registro de eventos", "registro.py", "Escribe JSON Lines con marca de tiempo de cada percepción, decisión, acción y fallo."],
], [4.2, 5.0, 7.2], fs=8, cap="Tabla 3. Correspondencia entre componentes y código")
P("Relación con los siete componentes IBM (los siete se verificaron en la sección 8 de este mismo documento):", bold=True)
T(["Componente IBM (sección 8)", "Implementación en el código", "Estado", "Limitación"], [
    ["1. Percepción y manejo de entradas", "percepcion.py", "Implementado", "Sensores simulados; el historial de pagos se reduce a dos variables numéricas."],
    ["2. Planificación y descomposición de tareas", "AgenteCreditScore.PLAN y evaluar()", "Parcial", "Secuencia fija de ocho tareas registrada en el log; no hay replanificación dinámica."],
    ["3. Memoria", "memoria.py", "Implementado", "En memoria RAM; guardar() persiste a archivo, pero no se carga automáticamente al iniciar."],
    ["4. Razonamiento y toma de decisiones", "decision.py", "Implementado", "Costos de decisión definidos por los autores, no por una política institucional real."],
    ["5. Llamada de herramientas y acción", "acciones.py y entorno.py", "Parcial", "Las «herramientas» son funciones de un entorno simulado; no hay APIs reales ni modelo de lenguaje."],
    ["6. Comunicación", "Payload a analista y originación; registro.py", "Parcial", "Salida estructurada con razón y tres factores principales; sin explicación en lenguaje natural."],
    ["7. Aprendizaje y adaptación", "modelo.py y revisar_aprendizaje()", "Implementado", "La revisión se invoca manualmente; no hay programación periódica."],
], [4.0, 4.0, 2.5, 5.9], fs=8, cap="Tabla 4. Componentes IBM y su implementación")
P("Ningún componente quedó sin implementar; los marcados como parciales tienen la limitación indicada. No se usa ningún modelo de lenguaje: todo el razonamiento lo realiza el código del repositorio.")
H("9.2 Tecnologías", 2)
T(["Tecnología", "Versión", "Función", "Justificación", "Alternativa considerada", "Motivo de la elección", "Uso"], [
    ["Python", "3.12.3", "Lenguaje del agente", "Permite el ciclo completo con poco código", "Java", "Mayor familiaridad y ecosistema científico", "Usada"],
    ["NumPy", "2.4.4", "Álgebra y regresión logística", "Única dependencia del agente", "scikit-learn 1.8.0", "Evitar dependencia y mostrar el mecanismo de aprendizaje (Newton-Raphson)", "Usada"],
    ["unittest (biblioteca estándar)", "Python 3.12.3", "Pruebas automatizadas", "No requiere instalación", "pytest", "pytest no estaba disponible sin red; las pruebas también corren con pytest", "Usada"],
    ["Graphviz", "2.43.0", "Diagrama de componentes", "Reproducible desde .dot", "draw.io", "Fuente en texto versionable", "Usada"],
    ["Mermaid (fuentes .mmd) y render_secuencia.py con Matplotlib 3.10.8", "—", "Diagramas de secuencia", "Fuente en texto; PNG generado por el mismo archivo", "PlantUML, mermaid-cli", "mermaid-cli no pudo ejecutarse (sin navegador); se escribió un renderizador propio", "Usada"],
    ["python-docx / LibreOffice", "1.2.0 / 24.2.7.2", "Generación del informe DOCX y PDF", "Informe reproducible desde datos reales", "Edición manual", "Evita transcribir cifras a mano", "Usada"],
    ["scikit-learn", "1.8.0", "Modelos de clasificación", "Biblioteca estándar de ML", "—", "Descartada para mantener una sola dependencia", "Solo evaluada"],
    ["Streamlit y pandas (app.py)", "1.66.0 (probada)", "Interfaz web y despliegue público (Opción A)", "Evalúa solicitudes y ejecuta las pruebas en vivo; despliegue gratuito desde GitHub", "Video (Opción B)", "Los integrantes no pueden grabar juntos; la interfaz no contiene lógica de decisión", "Usada"],
], [2.5, 1.5, 2.1, 2.6, 2.1, 3.9, 1.5], fs=7, cap="Tabla 5. Tecnologías usadas y evaluadas")
H("9.3 Entradas y sensores", 2)
T(["Entrada", "Tipo", "Rango o formato", "Origen", "Validaciones", "Ejemplo"], [
    ["id_solicitud", "Texto", "No vacío; existente en el entorno", "Interfaz / originación", "Si no existe: INVALIDA y revisión humana", "E01"],
    ["ingresos_mensuales", "Decimal (COP)", "[500.000; 50.000.000]", "Formulario (simulado)", "Numérico finito; faltante: INCOMPLETA; negativo o no numérico: INVALIDA; fuera de rango: FUERA_DOMINIO", "8.000.000"],
    ["obligaciones_mensuales", "Decimal (COP)", "[0; 100.000.000]", "Formulario (simulado)", "Igual que ingresos", "800.000"],
    ["dti (derivada)", "Decimal", "[0; 3] (obligaciones / ingresos)", "Calculada por el agente", "Fuera de rango: FUERA_DOMINIO (inconsistente)", "0,10"],
    ["pagos_atrasados_24m", "Entero", "[0; 24]", "Buró de crédito (simulado)", "Debe ser entero; fuera de rango: FUERA_DOMINIO; 3 reintentos si el buró falla", "0"],
    ["antiguedad_historial_meses", "Entero", "[0; 600]", "Buró de crédito (simulado)", "Igual que pagos atrasados", "120"],
    ["lineas_credito_abiertas", "Entero", "[0; 30]", "Buró de crédito (simulado)", "Igual que pagos atrasados", "3"],
    ["Casos históricos similares", "Vector de 5 variables y resultado 0/1", "Mínimo 25 casos con resultado", "Memoria de largo plazo", "Si hay menos de 25, no se aplica la regla R3", "3.500 casos de entrenamiento"],
    ["resultado_observado", "Entero 0/1", "0 = pagó; 1 = incumplió", "Cartera (simulada)", "Solo se asocia a solicitudes evaluadas", "1"],
], [3.5, 1.7, 2.7, 2.3, 4.3, 1.9], fs=7.5, cap="Tabla 6. Entradas del agente (config/config.json define los rangos)")
H("9.4 Salidas y acciones", 2)
aud = None
for linea in open(L / "escenarios_eventos.jsonl", encoding="utf-8"):
    e = json.loads(linea)
    if e["evento"] == "AUDITORIA" and e["id_solicitud"] == "E14":
        aud = e
T(["Salida / acción", "Formato", "Módulo", "Efecto esperado", "Ejemplo real (logs/escenarios_eventos.jsonl)"], [
    ["Aprobación o rechazo recomendado", "Diccionario: id, recomendacion, calificacion, p_default, version_modelo", "acciones.py · _originacion()", "Se registra en el sistema de originación; con 3 fallos queda DIFERIDA en el outbox", "E01: APROBAR, p = 0,0155, calificación A, REGISTRADA"],
    ["Remisión a revisión humana", "Diccionario: id, calificacion, p_default, razon, factores, reglas", "acciones.py · ejecutar()", "El caso entra en la bandeja del analista con la razón de la abstención", "E05: REVISION_HUMANA por R4 (DTI 0,8001)"],
    ["Calificación de riesgo", "Letra A–E por bandas de probabilidad (A < 5 %; B < 10 %; C < 20 %; D < 35 %; E ≥ 35 %)", "decision.py · calificar()", "Se registra en el expediente", "p = 0,0155 → A"],
    ["Registro histórico", "Línea JSON de auditoría con entradas, versión, probabilidad, acción y razón", "acciones.py · ejecutar(); registro.py", "Trazabilidad y datos para recalibración", "E14: " + (json.dumps({k: aud[k] for k in ("accion", "calificacion", "estado_acciones")}, ensure_ascii=False) if aud else "")],
], [3.0, 3.8, 2.8, 3.6, 3.2], fs=7.5, cap="Tabla 7. Salidas y actuadores")
H("9.5 Diagramas de secuencia UML", 2)
P("Los diagramas se escribieron en Mermaid (archivos .mmd) y se renderizaron con docs/diagramas/render_secuencia.py. Los mensajes siguen las llamadas reales del código; los fragmentos loop y alt corresponden a los ciclos de reintento y a las ramas de decisión. Las Figuras 2 a 4 aparecen en las páginas horizontales que siguen a los extractos de código.")
P("Explicación de la Figura 2. La interfaz invoca evaluar(id). La percepción lee el formulario y consulta el buró dentro de un ciclo de hasta tres intentos (mensajes 4 a 7). Con una percepción válida, la decisión obtiene la probabilidad y los factores del modelo (10 y 11) y la tasa de los casos similares (12 y 13). El fragmento alt de los mensajes 14 y 15 refleja el orden de las reglas: R4 (límite de DTI) y R3 (desacuerdo con casos similares) remiten a revisión; si no se activan, R2 elige la acción de menor costo esperado. Los actuadores registran la calificación, ejecutan la acción principal (alt de los mensajes 18 a 20) y registran auditoría y memoria (21). En el código, el evento AUDITORIA se escribe dentro de Actuadores.ejecutar(), antes del evento ACCION que emite el agente.")
P("Explicación de las Figuras 3 y 4. Con el buró caído, la percepción agota tres intentos, registra tres eventos SENSOR_FALLO y devuelve un estado INCOMPLETA; la regla R1 hace que el agente se abstenga y remita el caso, sin llamar a originación. Con la originación caída, la decisión sigue siendo APROBAR, pero tras tres intentos fallidos la recomendación se guarda en el outbox (estado DIFERIDA) y se audita, de modo que no se pierde. Ambos comportamientos están cubiertos por los escenarios E12 y E14 de las pruebas.")
P("Correspondencia con el código real (extractos del repositorio):", bold=True)
CODE(src(AgenteCreditScore.evaluar, "per = None", "estados = self.actuar", 12) + "\nestados = self.actuar(dec, per)            # ACCION", fs=7)
CODE(src(ModuloPercepcion._leer_buro, None, None, 10), fs=7)
CODE(src(ModuloDecision.decidir, "# R4", "# R2", 14), fs=7)
CODE(src(Actuadores._originacion, None, None, 12), fs=7)
pagina_landscape(True)
FIGFIT(D / "diagramas" / "secuencia_tipico.png", 26.0, 15.0, "Figura 2. Escenario típico: evaluación normal de una solicitud (fuente: docs/diagramas/secuencia_tipico.mmd)")
FIGFIT(D / "diagramas" / "secuencia_excepcion_buro.png", 26.0, 15.0, "Figura 3. Escenario de excepción 1: buró de crédito caído, escenario E12 (fuente: secuencia_excepcion_buro.mmd)", True)
FIGFIT(D / "diagramas" / "secuencia_excepcion_originacion.png", 26.0, 15.0, "Figura 4. Escenario de excepción 2: sistema de originación caído, escenario E14 (fuente: secuencia_excepcion_originacion.mmd)", True)
pagina_landscape(False)
H("9.6 Diferencias entre diseño e implementación", 2)
P("Se comparó cada elemento del diseño (secciones 2 a 8) con el código. El tipo de agente, el ambiente y el ciclo coinciden; las diferencias encontradas se listan a continuación.")
T(["Elemento del diseño", "Implementación realizada", "Diferencia encontrada", "Justificación", "Impacto"], [
    ["Agente de aprendizaje con rasgos de utilidad", "Regresión logística recalibrable y minimización del costo esperado", "Ninguna (coincide; ver pruebas M7 y escenarios E01 a E03)", "—", "—"],
    ["Ambiente parcialmente observable y estocástico", "Simulador con ruido latente no observable por el agente", "Ninguna", "—", "—"],
    ["Historial de pagos: datos tabulares, categóricos y temporales", "Dos variables numéricas: pagos atrasados en 24 meses y antigüedad", "Simplificación", "Tiempo disponible y ausencia de datos reales", "Menor riqueza predictiva"],
    ["Ingresos desde formulario, sistema financiero o API de verificación", "Solo formulario simulado", "Se omiten las otras fuentes", "No hay acceso a servicios reales", "Sin verificación cruzada de ingresos"],
    ["DTI calculado a partir de obligaciones registradas", "obligaciones_mensuales se agregó como entrada explícita", "Entrada nueva no listada en la Tabla 1", "Es necesaria para calcular el DTI", "Ninguno"],
    ["Historial de decisiones sobre solicitudes similares", "Búsqueda de los 25 casos similares con resultado conocido", "Se usan resultados, no las decisiones previas", "El resultado es la señal que permite contrastar el modelo", "Ninguno relevante"],
    ["Abstención ante «incertidumbre elevada»", "Regla R3: desacuerdo > 0,30 entre el modelo y los casos similares", "Se concretó la incertidumbre como un desacuerdo medible", "Debe ser operacionalizable", "Umbral 0,30 sin calibración empírica"],
    ["Políticas y umbrales de la entidad", "Costos 4 / 1 / 0,5 y DTI máximo 0,80 en config.json", "Valores definidos por los autores", "No hay política institucional real", "Cerca de la mitad de las solicitudes se remite a revisión"],
    ["Resultado observado llega por seguimiento de cartera", "El simulador lo entrega al instante", "Sin retraso temporal", "Simplificación de la simulación", "No se prueba el manejo de resultados tardíos"],
    ["Registro en BD de auditoría; API de originación; notificaciones", "Listas en memoria y archivo JSONL; funciones del entorno simulado", "Sin sistemas reales", "Sin acceso a infraestructura", "No se prueba integración real"],
    ["Recalibración periódica", "revisar_aprendizaje() invocable", "No está programada", "Alcance del curso", "Requiere activación manual"],
    ["Muestra de 5.000 solicitudes históricas reales", "5.000 solicitudes sintéticas (3.500 entrenamiento y 1.500 prueba)", "Datos sintéticos", "No se dispone de datos reales", "El AUC refleja la señal del simulador, no la de una entidad real"],
    ["Meta: reducir 20 % el tiempo frente al proceso manual", "Se mide la latencia p95 por evaluación; la reducción queda pendiente", "Medida reformulada", "No existe línea base manual", "Medida M2b pendiente"],
    ["Entrada por formulario", "app.py: interfaz web", "No estaba en el diseño", "Opción A", "Ninguno"],
    ["Métricas complementarias (precisión, recall, falsos positivos, subgrupos)", "ECE con umbral; recall y tasas como información sin umbral", "Subgrupos no evaluados", "El simulador no tiene atributos de subgrupo", "Equidad sin verificar"],
], [3.4, 3.8, 3.0, 3.2, 3.0], fs=7, cap="Tabla 8. Diferencias entre el diseño y la implementación")

# ============================ 9.7 PRUEBAS ============================
H("9.7 Pruebas de desempeño", 2)
P("Procedimiento. Los umbrales se fijaron en config/umbrales.json antes de ejecutar las pruebas (SHA-256 " + PR["meta"]["sha256_umbrales"][:16] + "…, registrado en logs/resultados_pruebas.json). Las pruebas con aleatoriedad se ejecutaron con 5 semillas (1 a 5) y se reportan media y desviación estándar. Los comandos para reproducirlas están en el README.")
T(["Medida", "Justificación del umbral (fijado antes de ejecutar)"], [
    ["M1. AUC-ROC", "0,75 es la meta del diseño y del caso de uso del catálogo."],
    ["M2. Latencia p95", "200 ms es un tiempo interactivo razonable para una recomendación; sustituye a la reducción del 20 % porque esta no tiene línea base."],
    ["M2b. Reducción del 20 %", "Meta original del diseño; pendiente de la línea base manual."],
    ["M3. Remisión ante entradas inválidas", "100 %: la validación es determinista y el diseño obliga a la abstención."],
    ["M4. ECE", "0,05 es un límite habitual de buena calibración con 10 bins."],
    ["M5. Trazabilidad", "100 %: el diseño exige registrar entradas, salida y versión de cada solicitud."],
    ["M6a. Desviación del AUC", "0,02 (dos puntos de AUC) como variación aceptable entre muestras."],
    ["M6b. Reproducibilidad", "100 %: con la misma semilla el agente debe ser idéntico."],
    ["M7a. Detección de deriva", "100 %: el diseño exige detectar cambios en la distribución."],
    ["M7b. Mejora tras recalibrar", "80 % de las semillas (4 de 5) como mejoría consistente."],
], [5.0, 11.4], fs=8, cap="Tabla 9. Umbrales y su justificación")
pagina_landscape(True)
H("Tabla de resultados obtenidos", 3)
filas = [r[:] for r in TABLA[1:]]
T(TABLA[0], filas, [5.0, 4.4, 1.8, 1.7, 4.9, 2.1, 6.0], fs=7, cap="Tabla 10. Resultados de las pruebas (fuente: logs/resultados_pruebas.csv; generado por scripts/run_pruebas_desempeno.py)")
P("Evidencia de funcionamiento (Opción A): la pestaña «3. Pruebas» de la app desplegada ejecuta en vivo las pruebas unitarias, los 17 escenarios y las medidas M1–M7 y muestra esta misma tabla; los archivos de la columna final son el respaldo.", italic=True, size=8)
pagina_landscape(False)
H("Escenarios de prueba deterministas (17 casos) y pruebas unitarias (20)", 3)
T(["ID", "Escenario", "Categoría", "Resultado esperado", "¿Cumple?"], [[e["id"], e["escenario"], e["categoria"], e["esperado"], "Sí" if e["cumple"] else "No"] for e in ES], [1.0, 5.6, 2.0, 6.6, 1.2], fs=7.5, cap="Tabla 11. Escenarios (fuente: logs/escenarios.csv y logs/escenarios_eventos.jsonl)")
P(f"Las {sum(e['cumple'] for e in ES)} de {len(ES)} pruebas de escenario cumplen lo esperado. Los 17 escenarios incluyen casos normales, de límite (DTI igual y superior a la política; 24 y 25 pagos atrasados), entradas inválidas, fallos esperados (buró caído, originación caída, solicitud inexistente) y condiciones adversas (fallos transitorios con recuperación); E17 repite el mismo caso 5 veces. Las 20 pruebas unitarias (logs/unittest.txt) pasan. Aparte, M3 ejecutó 1.000 entradas corrompidas de 10 tipos (100 por tipo) y M5 incluyó 500 evaluaciones con fallos aleatorios del 30 % en buró y originación.")
m = PR["M1_M2_M4_por_semilla"]
inf = f"Métricas informativas (sin umbral): % de solicitudes aprobadas {pct(min(x['pct_aprobar'] for x in m),0)}–{pct(max(x['pct_aprobar'] for x in m),0)}, rechazadas {pct(min(x['pct_rechazar'] for x in m),0)}–{pct(max(x['pct_rechazar'] for x in m),0)}, remitidas a revisión {pct(min(x['pct_revision'] for x in m),0)}–{pct(max(x['pct_revision'] for x in m),0)}; tasa de incumplimiento entre aprobadas {pct(min(x['default_entre_aprobados'] for x in m))}–{pct(max(x['default_entre_aprobados'] for x in m))} frente a {pct(min(x['tasa_default'] for x in m))}–{pct(max(x['tasa_default'] for x in m))} en la muestra; recall del incumplimiento (incumplidoras no aprobadas) {pct(min(x['recall_default_no_aprobado'] for x in m))}–{pct(max(x['recall_default_no_aprobado'] for x in m))}."
P(inf)
H("Análisis de resultados", 3)
m7 = PR["M7"]
dB = [x["brier_antes"] - x["brier_despues"] for x in m7]
ecd = [x["ece_despues"] for x in m7]
auc_s = [x["auc"] for x in m]
P(f"¿Qué se cumplió? Las nueve medidas ejecutables cumplen su umbral: AUC medio de {f2(sum(auc_s)/5,4)} (mínimo {f2(min(auc_s),4)}), p95 de latencia de {f2(PR['M2']['p95_ms'],3)} ms, remisión correcta de 1.000 de 1.000 entradas inválidas, ECE medio bajo 0,05, trazabilidad completa en 9.000 evaluaciones, reproducibilidad exacta y detección y mejora tras recalibración en 5 de 5 semillas. ¿Qué no se cumplió? Ninguna medida ejecutada incumplió su umbral; M2b (reducción del 20 %) está pendiente porque falta la línea base manual.")
P("Por qué estos resultados deben leerse con cautela (causas y evidencia):", bold=True)
B(f" el AUC sobre datos sintéticos refleja la señal que el propio simulador contiene; demuestra que la tubería funciona, no que el agente discrimine bien en una entidad real. Además, el margen es estrecho: la peor semilla obtuvo {f2(min(auc_s),4)} frente a un umbral de 0,75 (logs/resultados_pruebas.json#M1).", "M1:")
B(f" una latencia inferior a 1 ms se mide en un entorno simulado en memoria, sin red ni consulta real al buró; no es comparable con el proceso manual y por eso M2b sigue pendiente.", "M2:")
B(f" el agente remite a revisión humana entre {pct(min(x['pct_revision'] for x in m),0)} y {pct(max(x['pct_revision'] for x in m),0)} de las solicitudes válidas. Es consecuencia del costo de revisión (0,5) y de los costos de error (4 y 1) fijados por los autores; si esos costos no reflejan a la entidad, la carga de los analistas sería mayor de lo deseable.", "Carga de revisión:")
B(f" el PSI sobre las puntuaciones nunca superó su umbral de 0,2 (valores de {f2(min(x['psi'] for x in m7),3)} a {f2(max(x['psi'] for x in m7),3)}); la deriva se detectó por la brecha de calibración ({f2(min(x['brecha_calibracion'] for x in m7),3)} a {f2(max(x['brecha_calibracion'] for x in m7),3)} frente a 0,05). La deriva simulada cambia principalmente la tasa base de incumplimiento, que el PSI de las puntuaciones no captura bien. La recalibración mejoró el Brier en 5 de 5 semillas, pero la mejora es modesta ({f2(min(dB),4)} a {f2(max(dB),4)}), el ECE posterior sigue por encima de 0,05 en {sum(1 for x in ecd if x>0.05)} de 5 semillas (hasta {f2(max(ecd),3)}) y en la semilla 5 el AUC bajó de {f2(m7[4]['auc_antes'],4)} a {f2(m7[4]['auc_despues'],4)} (logs/resultados_pruebas.json#M7).", "M7:")
P("Cambios que podrían corregirlo: calcular el PSI por variable de entrada y monitorear la tasa de incumplimiento observada; recalibrar la probabilidad (por ejemplo, ajuste del intercepto o regresión isotónica) con más retroalimentación antes de reentrenar todo el modelo; calibrar los costos de decisión con datos de la entidad; obtener la línea base manual para cerrar M2b; y repetir M1 con datos reales o con un conjunto público de crédito. Los umbrales no se modificaron después de observar los resultados.")
H("9.8 Limitaciones y trabajo futuro", 2)
P("Limitaciones reales:", bold=True)
for t in ["Datos: todas las solicitudes son sintéticas y la señal predictiva la define el simulador; no hay evidencia de desempeño con datos reales ni de equidad entre subgrupos.",
          "Entorno: buró, originación y cartera son simulados; el resultado observado llega al instante, no meses después.",
          "Modelo: regresión logística con cinco variables y sin interacciones; la regla R3 usa un umbral (0,30) sin calibración empírica.",
          "Pruebas: la medida de reducción del 20 % del tiempo está pendiente; la latencia medida no incluye redes ni personas.",
          "Aprendizaje: la revisión de deriva debe invocarse manualmente y su detección por PSI resultó poco sensible al tipo de deriva simulada.",
          "Decisión: los costos de error y de revisión son supuestos de los autores."]:
    B(" " + t)
P("Mejoras propuestas, por prioridad:", bold=True)
for t in ["Alta: medir la línea base manual (mínimo 10 solicitudes) y cerrar M2b; los integrantes deben escribir el promedio en config/linea_base_manual.json.",
          "Alta: repetir M1 y M4 con un conjunto de datos público de riesgo crediticio o con datos reales anonimizados.",
          "Alta: añadir evaluación de equidad por subgrupos y revisión de variables sensibles.",
          "Media: PSI por variable y monitoreo de la tasa de incumplimiento; recalibración isotónica o por intercepto.",
          "Media: calibrar costos de decisión y el umbral de R3 con datos de la entidad.",
          "Baja: persistir la memoria de largo plazo entre ejecuciones, programar la revisión de aprendizaje."]:
    B(" " + t)
H("9.9 Enlaces", 2)
T(["Recurso", "Estado", "Cómo obtenerlo"], [
    ["Repositorio (GitHub)", "Publicado (comprobado por el integrante: abre sin iniciar sesión)", URL_REPO],
    ["Plataforma de despliegue (Opción A: Streamlit Community Cloud)",
     "PENDIENTE: no se ha desplegado" if URL_APP == "PENDIENTE" else "Publicada; disponible hasta la fecha de calificación",
     "Desplegar siguiendo docs/guia_despliegue.md y pegar aquí la URL." if URL_APP == "PENDIENTE" else URL_APP],
    ["Instrucciones de uso y casos de prueba sugeridos", "En el repositorio", "docs/guia_despliegue.md y pestaña «Evaluar una solicitud» de la app"],
    ["Video de la demostración", "No aplica: se eligió la Opción A (plataforma de despliegue)", "—"],
], [4.5, 4.5, 7.4], fs=8, cap="Tabla 12. Enlaces de entrega")
PB([("Declaración de uso de herramientas de IA. ", True), ("El código, los diagramas y este documento se prepararon con ayuda de un asistente de IA (Claude, de Anthropic). Los resultados numéricos provienen de ejecuciones reales del código incluido. La bitácora docs/bitacora_uso_ia.md registra el uso; los campos de verificación humana están pendientes de completar por los integrantes, que deben revisar y poder explicar el código y el documento.", False)])

# ============================ 10 CONCLUSIONES ============================
H("10. Conclusiones", 1)
P("El diseño del CREDITSCORE-AGENT se implementó como un agente de aprendizaje con decisión por utilidad, con un ciclo Percepción → Decisión → Acción identificable en el código, memoria de corto y largo plazo, abstención ante datos inválidos o incertidumbre, actuadores con manejo de fallos y registro con marca de tiempo. Las pruebas ejecutadas muestran que el agente cumple los umbrales definidos para discriminación, calibración, tiempo de respuesta, tolerancia a entradas inválidas, trazabilidad, reproducibilidad y adaptación a la deriva, siempre sobre datos sintéticos.")
P("Estos resultados no demuestran todavía que el agente sea apto para una entidad financiera. Falta la línea base manual para la medida de tiempo, la validación con datos reales, la evaluación de equidad por subgrupos y la calibración de los costos de decisión con la política de una entidad. La decisión de implementar este agente frente a Cartrescue-agent se apoyó en puntajes estimados para este último y debe confirmarse con su diseño real. Por el impacto de las decisiones crediticias, el agente debe seguir diseñándose como apoyo a la decisión, con revisión humana, auditoría y control de sesgos.")
H("Referencias", 1)
for r in ["Curso de Inteligencia Artificial. (2026). Intelligent Agents: Use Cases [Documento de casos de uso]. Material proporcionado en la asignatura.",
          "Harris, C. R., Millman, K. J., van der Walt, S. J., et al. (2020). Array programming with NumPy. Nature, 585, 357–362. https://doi.org/10.1038/s41586-020-2649-2",
          "IBM. (2025). Componentes de los agentes de IA. IBM Think. https://www.ibm.com/mxes/think/topics/components-of-ai-agents (enlace por verificar).",
          "Mota, J., y Moncayo, C. (2023). Un análisis del perfil de riesgo en la dinámica de préstamos persona a persona mediante clústeres de K-medias. Economía: Teoría y Práctica, 31(59), 119–138.",
          "Nortey, E. N. N., Asante-Koranteng, J., Atemkeng, M., Ansah-Narh, T., Mensah, D., Davis, R., y Laryea, R. A. (2026). An Optimised Greedy-Weighted Ensemble Framework for Financial Loan Default Prediction. arXiv:2603.18927 (referencia heredada de la entrega anterior; los integrantes deben verificar que exista antes de entregar).",
          "Russell, S. J., y Norvig, P. (2021). Artificial Intelligence: A Modern Approach (4.ª ed. global). Pearson. https://aima.cs.berkeley.edu/globalindex.html"]:
    P(r, size=9)
out = D / "Documento_Diseno_Ajustado.docx"
doc.save(out); print("guardado", out)
